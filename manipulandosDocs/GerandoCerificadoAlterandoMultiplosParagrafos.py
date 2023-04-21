from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook

# abrindo o arquivo
nomeArquivoAlunos = "DadosAlunos.xlsx"
planilhaDadosAlunos = load_workbook(nomeArquivoAlunos)

# selecionando a planilha
sheetSelecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2, len(sheetSelecionada["A"])+1):
    # abrindo arquivo já existente
    arquivoWord = Document("C:\\Users\\murilo.fernandes\\Documents\\Udemy\\Robotic_Process_Automation_RPA_automação_Python\\manipulandosDocs\\Certificado3.docx")

    # seleiconando estilo
    estilo = arquivoWord.styles["Normal"]
    # pegando o nome do aluno quando passar em cada célula
    nomeAluno = sheetSelecionada['A%s' % linha].value
    dia = sheetSelecionada['B%s' % linha].value
    mes = sheetSelecionada['C%s' % linha].value
    ano = sheetSelecionada['D%s' % linha].value
    curso = sheetSelecionada['E%s' % linha].value
    instrutor = sheetSelecionada['F%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calabri (Corpo)"
            fonte.size = Pt(24)

        paragrafoParte1 = "Concluiu com sucesso o curso de "
        paragrafoParte2 = ", como carga horária de 20 horas, promovido pela escola de Cursos Online em"
        paragrafoCompleto = f"{paragrafoParte1} {curso}{paragrafoParte2} {dia} de {mes} de {ano}"

        if "escola" in paragrafo.text:
            paragrafo.text = paragrafoCompleto
            fonte = estilo.font
            fonte.name = "Calabri (Corpo)"
            fonte.size = Pt(24)

        if "instrutor" in paragrafo.text:
            paragrafo.text = instrutor + " - Instrutor"
            fonte = estilo.font
            fonte.name = "Calabri (Corpo)"
            fonte.size = Pt(24)

    # definindo o caminho da pasta para salvar os certificados
    novoCaminho = f"C:\\Users\\murilo.fernandes\\Documents\\Udemy\\Robotic_Process_Automation_RPA_automação_Python\\manipulandosDocs\\certificadosAlunos\\{nomeAluno}2.docx"
    arquivoWord.save(novoCaminho)

print("Certificados gerados com sucesso!")
