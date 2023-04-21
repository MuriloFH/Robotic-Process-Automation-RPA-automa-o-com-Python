from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os

# abrindo o arquivo
nomeArquivoAlunos = "Alunos.xlsx"
planilhaDadosAlunos = load_workbook(nomeArquivoAlunos)

# selecionando a planilha
sheetSelecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2, len(sheetSelecionada["A"])+1):
    # abrindo arquivo já existente
    arquivoWord = Document("C:\\Users\\murilo.fernandes\\Documents\\Udemy\\Robotic_Process_Automation_RPA_automação_Python\\manipulandosDocs\\Certificado1.docx")

    # seleiconando estilo
    estilo = arquivoWord.styles["Normal"]
    # pegando o nome do aluno quando passar na célula A
    nomeAluno = sheetSelecionada['A%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calabri (Corpo)"
            fonte.size = Pt(24)

    # definindo o caminho da pasta para salvar os certificados
    novoCaminho = f"C:\\Users\\murilo.fernandes\\Documents\\Udemy\\Robotic_Process_Automation_RPA_automação_Python\\manipulandosDocs\\certificadosAlunos\\{nomeAluno}.docx"
    arquivoWord.save(novoCaminho)

print("Certificados gerados com sucesso!")

